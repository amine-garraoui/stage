"""
reports/excel.py - Export mensuel au format du classeur officiel ASKit.

L'export privilegie le classeur modele:
  C:\\Users\\MSI\\Downloads\\PROJETSTAGE\\ASKit - Report request enregistres - Juin 2026.xlsx

Le modele conserve la mise en page et les graphiques. Le code remplace uniquement:
  - les cellules sources de la feuille TDB,
  - les lignes des feuilles tickets ouverts, tickets fermes et enquete.
"""

from __future__ import annotations

from copy import copy
import datetime as dt
import os
from pathlib import Path

import pandas as pd
from openpyxl import Workbook, load_workbook
from openpyxl.styles import Alignment, Font, PatternFill

from core import kpi


TEMPLATE_ENV = "ASKIT_TDB_TEMPLATE"
TEMPLATE_GLOB = "ASKit - Report request enregistr* - Juin 2026.xlsx"
SHEET_TDB = "TDB"
SHEET_OUVERTS = "tickets ouverts"
SHEET_FERMES = "tickets fermés"
SHEET_ENQUETE = "Askit -SST KRM enquête"

OFFICIAL_SITES = ["Tunisie/Megrine", "Tunisie/Kram", "Tunisie/Sousse"]
OFFICIAL_SITE_POLES = [
    ("Tunisie/Megrine", "AVS"),
    ("Tunisie/Megrine", "BBS"),
    ("Tunisie/Megrine", "E&T"),
    ("Tunisie/Megrine", "supports"),
    ("Tunisie/Kram", "E&T"),
    ("Tunisie/Kram", "supports"),
    ("Tunisie/Kram", "BBS"),
    ("Tunisie/Sousse", "E&T"),
    ("Tunisie/Sousse", "supports"),
]

OPEN_HEADERS = [
    "N° ticket", "Bénéficiaire", "Bénéficiaire ID", "Directeur", "département ",
    "Bénéficiaire : Localisation", "Enregistré le", "Date de résolution", "Sujet",
    "Priorité", "E_REPONSES", "Description", "Bénéficiaire : Niveau de VIP",
    "GROUP_$lng", "Groupe courant", "Résolu par (groupe)",
    "Résolu par (intervenant)", "GROUP_$lng", "Meta Statut",
    "Délai de résolution (min)", "Statut ticket", "Workflow",
]

CLOSED_HEADERS = [
    "N° ticket", "Bénéficiaire", "Bénéficiaire ID", "Enregistré le",
    "Date de résolution", "Sujet", "Priorité", "E_REPONSES", "Description",
    "Bénéficiaire : Niveau de VIP", "Bénéficiaire : Localisation", "GROUP_$lng",
    "Groupe courant", "Résolu par (groupe)", "Résolu par (intervenant)",
    "GROUP_$lng", "Meta Statut", "Délai de résolution (min)",
    "Statut ticket", "Workflow",
]

ENQUETE_HEADERS = [
    "N° ticket", "Date de création", "Résolu par (groupe)", "Date de résolution",
    "Résolu par (intervenant)", "Sujet", "Satisfaction de traitement",
    "Communication des opérateurs", "Satisfaction du temps de traitement",
    "Commentaire enquête",
]


def _find_sheet(wb, desired: str):
    normalized = _normalize(desired)
    for ws in wb.worksheets:
        if _normalize(ws.title) == normalized:
            return ws
    return wb[desired]


def _normalize(value: object) -> str:
    import unicodedata

    text = str(value).strip().lower()
    text = unicodedata.normalize("NFKD", text)
    text = "".join(ch for ch in text if not unicodedata.combining(ch))
    return " ".join(text.split())


def _template_path() -> Path | None:
    configured = os.environ.get(TEMPLATE_ENV)
    if configured:
        path = Path(configured).expanduser()
        if path.exists():
            return path

    repo_template_dir = Path(__file__).resolve().parents[1] / "assets" / "templates"
    if repo_template_dir.exists():
        matches = sorted(repo_template_dir.glob(TEMPLATE_GLOB))
        if matches:
            return matches[0]

    downloads_project = Path.home() / "Downloads" / "PROJETSTAGE"
    if downloads_project.exists():
        matches = sorted(downloads_project.glob(TEMPLATE_GLOB))
        if matches:
            return matches[0]

    return None


def _month_start(mois: str):
    try:
        return pd.Period(mois, "M").start_time.to_pydatetime()
    except Exception:
        return mois


def _clean(value):
    if value is None or value is pd.NA:
        return None
    try:
        if pd.isna(value):
            return None
    except (TypeError, ValueError):
        pass
    if isinstance(value, pd.Timestamp):
        return value.to_pydatetime()
    if isinstance(value, pd.Timedelta):
        return value.to_pytimedelta()
    return value


def _minutes_to_timedelta(value) -> dt.timedelta | None:
    try:
        if value is None or pd.isna(value):
            return None
        return dt.timedelta(minutes=float(value))
    except (TypeError, ValueError):
        return None


def _cell_style_from(src, dst) -> None:
    dst.font = copy(src.font)
    dst.fill = copy(src.fill)
    dst.border = copy(src.border)
    dst.alignment = copy(src.alignment)
    dst.number_format = src.number_format
    dst.protection = copy(src.protection)


def _replace_sheet(ws, data: pd.DataFrame, headers: list[str]) -> None:
    for col_idx, header in enumerate(headers, start=1):
        ws.cell(1, col_idx).value = header

    max_cols = len(headers)
    if ws.max_row > 1:
        for row in ws.iter_rows(min_row=2, max_row=ws.max_row, max_col=max_cols):
            for cell in row:
                cell.value = None

    template_row = [ws.cell(2, col_idx) for col_idx in range(1, max_cols + 1)] if ws.max_row >= 2 else []
    for row_idx, (_, row) in enumerate(data.iterrows(), start=2):
        values = list(row.values)
        for col_idx, header in enumerate(headers, start=1):
            cell = ws.cell(row_idx, col_idx)
            if template_row:
                _cell_style_from(template_row[col_idx - 1], cell)
            cell.value = _clean(values[col_idx - 1] if col_idx <= len(values) else None)


def _site_count(df: pd.DataFrame, mois: str, site: str) -> int:
    data = kpi.tickets_by_site(df, mois)
    if data.empty:
        return 0
    matched = data[data["site"].astype(str).str.casefold().eq(site.casefold())]
    return int(matched["nombre_tickets"].sum()) if not matched.empty else 0


def _site_pole_count(df: pd.DataFrame, mois: str, site: str, pole: str) -> int:
    data = kpi.tickets_by_site_pole(df, mois)
    if data.empty:
        return 0
    matched = data[
        data["site"].astype(str).str.casefold().eq(site.casefold())
        & data["pole"].astype(str).str.casefold().eq(pole.casefold())
    ]
    return int(matched["nombre_tickets"].sum()) if not matched.empty else 0


def _subject_counts_first_seen(df: pd.DataFrame, mois: str, limit: int = 40) -> list[tuple[str, int]]:
    opened = kpi.filter_opened_month(df, mois)
    if opened.empty or "categorie" not in opened.columns:
        return []
    values = opened["categorie"].fillna("Donnee indisponible").astype(str)
    counts = values.value_counts(sort=False)
    return [(subject, int(counts[subject])) for subject in counts.index[:limit]]


def _detail_counts(detail) -> list[int]:
    if detail is None:
        return [0, 0, 0, 0, 0]
    return [
        detail.tres_insatisfait,
        detail.plutot_insatisfait,
        detail.insatisfait,
        detail.satisfait,
        detail.tres_satisfait,
    ]


def _build_tickets_ouverts(df: pd.DataFrame, mois: str) -> pd.DataFrame:
    opened = kpi.filter_opened_month(df, mois).copy()
    rows = []
    for _, row in opened.iterrows():
        delay = _minutes_to_timedelta(row.get("delai_resolution_minutes"))
        rows.append({
            "N° ticket": row.get("ticket_id"),
            "Bénéficiaire": row.get("beneficiaire", ""),
            "Bénéficiaire ID": row.get("beneficiaire_id", ""),
            "Directeur": row.get("directeur", ""),
            "département ": row.get("pole", ""),
            "Bénéficiaire : Localisation": row.get("site", ""),
            "Enregistré le": row.get("date_ouverture"),
            "Date de résolution": row.get("date_fermeture"),
            "Sujet": row.get("categorie", ""),
            "Priorité": row.get("priorite"),
            "E_REPONSES": row.get("e_reponses", ""),
            "Description": row.get("description", ""),
            "Bénéficiaire : Niveau de VIP": row.get("vip_level", ""),
            "GROUP_$lng": row.get("group_lng", ""),
            "Groupe courant": row.get("groupe_traitant", ""),
            "Résolu par (groupe)": row.get("groupe_resolu", ""),
            "Résolu par (intervenant)": row.get("intervenant", ""),
            "Meta Statut": row.get("statut", ""),
            "Délai de résolution (min)": delay,
            "Statut ticket": row.get("statut_detail", ""),
            "Workflow": row.get("workflow", ""),
        })
    return pd.DataFrame(rows, columns=OPEN_HEADERS)


def _build_tickets_fermes(df: pd.DataFrame, mois: str) -> pd.DataFrame:
    closed = kpi.filter_closed_month(df, mois).copy()
    rows = []
    for _, row in closed.iterrows():
        delay = _minutes_to_timedelta(row.get("delai_resolution_minutes"))
        rows.append({
            "N° ticket": row.get("ticket_id"),
            "Bénéficiaire": row.get("beneficiaire", ""),
            "Bénéficiaire ID": row.get("beneficiaire_id", ""),
            "Enregistré le": row.get("date_ouverture"),
            "Date de résolution": row.get("date_fermeture"),
            "Sujet": row.get("categorie", ""),
            "Priorité": row.get("priorite"),
            "E_REPONSES": row.get("e_reponses", ""),
            "Description": row.get("description", ""),
            "Bénéficiaire : Niveau de VIP": row.get("vip_level", ""),
            "Bénéficiaire : Localisation": row.get("site", ""),
            "GROUP_$lng": row.get("group_lng", ""),
            "Groupe courant": row.get("groupe_traitant", ""),
            "Résolu par (groupe)": row.get("groupe_resolu", ""),
            "Résolu par (intervenant)": row.get("intervenant", ""),
            "Meta Statut": row.get("statut", ""),
            "Délai de résolution (min)": delay,
            "Statut ticket": row.get("statut_detail", ""),
            "Workflow": row.get("workflow", ""),
        })
    return pd.DataFrame(rows, columns=CLOSED_HEADERS)


def _build_enquete(df: pd.DataFrame, mois: str) -> pd.DataFrame:
    responses = df.attrs.get("satisfaction_responses")
    if responses is None or responses.empty:
        return pd.DataFrame(columns=ENQUETE_HEADERS)

    if "mois_enquete" in responses.columns and responses["mois_enquete"].notna().any():
        scoped = responses[responses["mois_enquete"].eq(mois)].copy()
        if scoped.empty:
            scoped = responses.iloc[0:0].copy()
    else:
        scoped = responses.copy()

    rows = []
    for _, row in scoped.iterrows():
        rows.append({
            "N° ticket": row.get("ticket_id"),
            "Date de création": row.get("date_enquete"),
            "Résolu par (groupe)": row.get("groupe_resolu", ""),
            "Date de résolution": row.get("date_resolution"),
            "Résolu par (intervenant)": row.get("intervenant", ""),
            "Sujet": row.get("categorie", ""),
            "Satisfaction de traitement": row.get("satisfaction_traitement"),
            "Communication des opérateurs": row.get("communication_operateurs"),
            "Satisfaction du temps de traitement": row.get("satisfaction_temps"),
            "Commentaire enquête": row.get("commentaire", ""),
        })
    return pd.DataFrame(rows, columns=ENQUETE_HEADERS)


def _write_tdb_values(wb, df: pd.DataFrame, mois: str) -> None:
    ws = _find_sheet(wb, SHEET_TDB)
    result = kpi.calculate_month_kpi(df, mois)
    sat = result.satisfaction
    month_value = _month_start(mois)

    for cell_ref in ["C2", "C19", "C24"]:
        ws[cell_ref] = month_value

    ws["C3"] = result.total_ouverts
    ws["C4"] = result.ouverts_et_fermes_meme_mois or 0
    ws["C5"] = result.total_fermes or 0
    ws["C6"] = _minutes_to_timedelta(result.delai_moyen_minutes)
    ws["C6"].number_format = "[h]:mm:ss"

    for row_idx, site in zip(range(20, 23), OFFICIAL_SITES):
        ws.cell(row_idx, 3).value = _site_count(df, mois, site)

    for row_idx, (site, pole) in zip(range(25, 34), OFFICIAL_SITE_POLES):
        ws.cell(row_idx, 3).value = _site_pole_count(df, mois, site, pole)

    subjects = _subject_counts_first_seen(df, mois)
    for row_idx in range(48, 88):
        offset = row_idx - 48
        if offset < len(subjects):
            subject, count = subjects[offset]
            ws.cell(row_idx, 2).value = subject
            ws.cell(row_idx, 3).value = count
        else:
            ws.cell(row_idx, 2).value = None
            ws.cell(row_idx, 3).value = None

    detail_rows = [
        _detail_counts(sat.detail_traitement),
        _detail_counts(sat.detail_communication),
        _detail_counts(sat.detail_temps),
    ]
    for row_idx, values in zip(range(92, 95), detail_rows):
        for col_idx, value in zip(range(3, 8), values):
            ws.cell(row_idx, col_idx).value = value

    total_notes = sum(sum(row) for row in detail_rows)
    for col_idx in range(3, 8):
        level_total = sum(ws.cell(row_idx, col_idx).value or 0 for row_idx in range(92, 95))
        cell = ws.cell(95, col_idx)
        cell.value = (level_total / total_notes) if total_notes else 0
        cell.number_format = "0.00%"

    ws["C98"] = result.total_fermes or 0
    ws["C99"] = sat.nombre_reponses
    ws["C100"] = (sat.taux_participation or 0) / 100
    ws["C100"].number_format = "0.00%"

    try:
        wb.calculation.fullCalcOnLoad = True
        wb.calculation.forceFullCalc = True
    except AttributeError:
        pass


def _new_fallback_workbook() -> Workbook:
    wb = Workbook()
    wb.active.title = SHEET_TDB
    wb.create_sheet(SHEET_OUVERTS)
    wb.create_sheet(SHEET_FERMES)
    wb.create_sheet(SHEET_ENQUETE)
    for ws in wb.worksheets:
        ws.freeze_panes = "A2"
        for cell in ws[1]:
            cell.font = Font(bold=True, color="FFFFFF")
            cell.fill = PatternFill("solid", fgColor="2E75B6")
            cell.alignment = Alignment(horizontal="center")
    return wb


def _load_workbook_from_template() -> Workbook:
    template = _template_path()
    if template is None:
        return _new_fallback_workbook()
    return load_workbook(template)


def export_month_excel(df: pd.DataFrame, mois: str, output_dir: str | Path) -> Path:
    """
    Genere le rapport Excel mensuel avec la meme structure que le fichier officiel.

    Les fichiers optionnels absents ne bloquent pas l'export: les feuilles et cellules
    correspondantes restent vides ou a zero.
    """
    output = Path(output_dir)
    output.mkdir(parents=True, exist_ok=True)
    path = output / f"ASKit - Report request enregistrés - {mois}.xlsx"

    wb = _load_workbook_from_template()

    ws_ouverts = _find_sheet(wb, SHEET_OUVERTS)
    ws_fermes = _find_sheet(wb, SHEET_FERMES)
    ws_enquete = _find_sheet(wb, SHEET_ENQUETE)

    _replace_sheet(ws_ouverts, _build_tickets_ouverts(df, mois), OPEN_HEADERS)
    _replace_sheet(ws_fermes, _build_tickets_fermes(df, mois), CLOSED_HEADERS)
    _replace_sheet(ws_enquete, _build_enquete(df, mois), ENQUETE_HEADERS)
    _write_tdb_values(wb, df, mois)

    try:
        wb.save(path)
        return path
    except PermissionError:
        timestamp = dt.datetime.now().strftime("%Y%m%d_%H%M%S")
        fallback = output / f"ASKit - Report request enregistrés - {mois}_{timestamp}.xlsx"
        wb.save(fallback)
        return fallback
