"""Excel Report Generator — openpyxl is installed."""

from __future__ import annotations

import io
from datetime import datetime
from typing import Optional

import pandas as pd
from openpyxl import Workbook
from openpyxl.styles import Alignment, Border, Font, PatternFill, Side
from openpyxl.utils import get_column_letter

PRIMARY = "003366"
ACCENT = "0066CC"
HEADER_FONT = Font(name="Calibri", bold=True, color="FFFFFF", size=11)
HEADER_FILL = PatternFill("solid", fgColor=PRIMARY)
ACCENT_FILL = PatternFill("solid", fgColor=ACCENT)
ALT_FILL = PatternFill("solid", fgColor="EEF2F8")
NORMAL = Font(name="Calibri", size=10)
BOLD = Font(name="Calibri", bold=True, size=10)
THIN = Side(border_style="thin", color="CCCCCC")
BORDER = Border(left=THIN, right=THIN, top=THIN, bottom=THIN)
CENTER = Alignment(horizontal="center", vertical="center", wrap_text=True)
LEFT = Alignment(horizontal="left", vertical="center", wrap_text=True)


def _h(ws, r, c, v, fill=None):
    cell = ws.cell(row=r, column=c, value=v)
    cell.font = HEADER_FONT; cell.fill = fill or HEADER_FILL
    cell.alignment = CENTER; cell.border = BORDER


def _d(ws, r, c, v, bold=False):
    cell = ws.cell(row=r, column=c, value=v)
    cell.font = BOLD if bold else NORMAL
    cell.alignment = LEFT; cell.border = BORDER
    if r % 2 == 0:
        cell.fill = ALT_FILL


def _widths(ws, widths):
    for i, w in enumerate(widths, 1):
        ws.column_dimensions[get_column_letter(i)].width = w


class ExcelReportGenerator:
    def generate(self, kpi_df: pd.DataFrame, insight_data: Optional[dict] = None,
                 anomaly_df: Optional[pd.DataFrame] = None,
                 forecast_df: Optional[pd.DataFrame] = None) -> bytes:
        wb = Workbook()
        wb.remove(wb.active)
        self._summary(wb, kpi_df, insight_data)
        self._history(wb, kpi_df)
        self._sites(wb, kpi_df)
        self._topics(wb, kpi_df)
        if anomaly_df is not None and not anomaly_df.empty:
            self._anomalies(wb, anomaly_df)
        if forecast_df is not None and not forecast_df.empty:
            self._forecasts(wb, forecast_df)
        buf = io.BytesIO()
        wb.save(buf)
        return buf.getvalue()

    def _summary(self, wb, kpi_df, insight):
        ws = wb.create_sheet("Résumé Exécutif")
        ws.sheet_view.showGridLines = False
        ws.merge_cells("A1:F1")
        c = ws["A1"]
        c.value = "RSI SAGEMCOM — Rapport de Performance IT Support"
        c.font = Font(name="Calibri", bold=True, size=16, color=PRIMARY)
        c.alignment = CENTER
        ws.row_dimensions[1].height = 35
        ws.merge_cells("A2:F2")
        ws["A2"].value = f"Généré le {datetime.now().strftime('%d/%m/%Y %H:%M')} — CONFIDENTIEL"
        ws["A2"].font = Font(name="Calibri", size=10, color="CC3300", italic=True)
        ws["A2"].alignment = CENTER
        if not kpi_df.empty:
            latest = kpi_df.iloc[-1]
            _h(ws, 4, 1, "Indicateur"); _h(ws, 4, 2, "Valeur")
            items = [
                ("Période", latest.get("period", "")),
                ("Tickets Ouverts", int(latest.get("tickets_opened", 0) or 0)),
                ("Tickets Fermés", int(latest.get("tickets_closed", 0) or 0)),
                ("En Cours EOM", int(latest.get("tickets_open_eom", 0) or 0)),
                ("Délai Moy. Résolution", f"{latest['avg_resolution_hours']:.1f}h"
                 if latest.get("avg_resolution_hours") else "N/D"),
                ("Satisfaction Moy.", f"{latest['avg_satisfaction']:.2f}/5"
                 if latest.get("avg_satisfaction") else "N/D"),
                ("Taux Participation", f"{latest['participation_rate']:.1f}%"
                 if latest.get("participation_rate") else "N/D"),
                ("Indice Performance RSI", f"{latest['performance_index']:.1f}/100"
                 if latest.get("performance_index") else "N/D"),
            ]
            for i, (lbl, val) in enumerate(items, 5):
                _d(ws, i, 1, lbl, bold=True); _d(ws, i, 2, val)
        if insight:
            ws.cell(row=14, column=1, value="Résumé").font = BOLD
            ws.merge_cells("A15:F20")
            c = ws["A15"]
            c.value = insight.get("summary_text", "")
            c.alignment = Alignment(horizontal="left", vertical="top", wrap_text=True)
        _widths(ws, [30, 25, 15, 15, 15, 15])

    def _history(self, wb, df):
        ws = wb.create_sheet("Historique KPI")
        ws.freeze_panes = "A2"
        cols = [("Période","period"),("Ouverts","tickets_opened"),("Fermés","tickets_closed"),
                ("Ouverts+Fermés","tickets_opened_and_closed"),("EOM","tickets_open_eom"),
                ("Délai Moy.h","avg_resolution_hours"),("Délai Médian h","median_resolution_hours"),
                ("Délai P90 h","p90_resolution_hours"),("Satisfaction","avg_satisfaction"),
                ("Participation %","participation_rate"),("Perf. Index","performance_index")]
        for ci, (lbl, _) in enumerate(cols, 1):
            _h(ws, 1, ci, lbl)
        for ri, (_, row) in enumerate(df.iterrows(), 2):
            for ci, (_, fld) in enumerate(cols, 1):
                val = row.get(fld)
                if isinstance(val, float) and pd.isna(val):
                    val = None
                elif isinstance(val, float):
                    val = round(val, 2)
                _d(ws, ri, ci, val)
        _widths(ws, [12,12,12,18,10,12,14,12,12,14,12])

    def _sites(self, wb, df):
        import json
        ws = wb.create_sheet("Par Site")
        _h(ws, 1, 1, "Période"); _h(ws, 1, 2, "Site"); _h(ws, 1, 3, "Tickets")
        ri = 2
        for _, row in df.iterrows():
            sites = row.get("by_site") or {}
            if isinstance(sites, str):
                try: sites = json.loads(sites)
                except: sites = {}
            for site, cnt in sorted(sites.items(), key=lambda x: -x[1]):
                _d(ws, ri, 1, row.get("period")); _d(ws, ri, 2, site); _d(ws, ri, 3, int(cnt))
                ri += 1
        _widths(ws, [12, 30, 12])

    def _topics(self, wb, df):
        import json
        ws = wb.create_sheet("Par Sujet")
        _h(ws, 1, 1, "Période"); _h(ws, 1, 2, "Sujet"); _h(ws, 1, 3, "Tickets")
        ri = 2
        for _, row in df.iterrows():
            topics = row.get("by_topic") or {}
            if isinstance(topics, str):
                try: topics = json.loads(topics)
                except: topics = {}
            for topic, cnt in sorted(topics.items(), key=lambda x: -x[1]):
                _d(ws, ri, 1, row.get("period")); _d(ws, ri, 2, topic); _d(ws, ri, 3, int(cnt))
                ri += 1
        _widths(ws, [12, 40, 12])

    def _anomalies(self, wb, df):
        ws = wb.create_sheet("Anomalies")
        hdrs = ["Période","Type","Sévérité","Portée","Observé","Référence","Écart %","Explication"]
        for ci, h in enumerate(hdrs, 1): _h(ws, 1, ci, h)
        for ri, (_, row) in enumerate(df.iterrows(), 2):
            vals = [f"{int(row.get('year',0)):04d}-{int(row.get('month',0)):02d}",
                    row.get("anomaly_type",""), row.get("severity",""), row.get("scope",""),
                    row.get("observed_value"), row.get("baseline_value"),
                    row.get("deviation_percent"), row.get("explanation","")]
            for ci, v in enumerate(vals, 1): _d(ws, ri, ci, v)
        _widths(ws, [12,18,12,20,12,12,10,60])

    def _forecasts(self, wb, df):
        ws = wb.create_sheet("Prévisions")
        hdrs = ["Période","Métrique","Valeur","Borne Inf.","Borne Sup.","Modèle","Historique"]
        for ci, h in enumerate(hdrs, 1): _h(ws, 1, ci, h)
        for ri, (_, row) in enumerate(df.iterrows(), 2):
            vals = [row.get("period",""), row.get("metric",""), row.get("value"),
                    row.get("lower"), row.get("upper"), row.get("model",""),
                    "Oui" if row.get("is_historical") else "Non"]
            for ci, v in enumerate(vals, 1): _d(ws, ri, ci, v)
        _widths(ws, [12,25,14,12,12,22,10])
